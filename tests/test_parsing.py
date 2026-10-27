"""Título, cabeçalho multinível, stub e territórios."""
import pytest

from pense.ingestion.readers import read_csv, read_xlsx
from pense.parsing.header import Vocab, parse_header
from pense.parsing.stub import classify_group_label, parse_stub
from pense.parsing.title import extract_title, parse_title, period_code
from pense.standardization.territories import TerritoryResolver

T2024 = ("Tabela 17.2.1 - Percentual de escolares de 13 a 17 anos em escolas cujo diretor(a) ou responsável "
         "tiveram de suspender ou interromper as aulas por motivo de segurança em termos de violência alguma vez "
         "nos 12 meses anteriores à pesquisa, e distribuição percentual por número de vezes, com indicação do "
         "intervalo de confiança de 95%, segundo os grupos de idade, a dependência administrativa da escola e as "
         "Grandes Regiões - 2024")
T2009 = ("Tabela 8.8 - Percentual de escolares frequentando o 9º ano menores de 18 anos de idade que, nos últimos "
         "30 dias, dirigiram veículo motorizado, por sexo e dependência administrativa da escola, com indicação do "
         "intervalo de confiança de 95%, segundo os municípios das capitais e Distrito Federal - 2009")
T2015 = ("Tabela 1.2.6.2 - Estimativas total e percentual de escolares frequentando o 9º ano do ensino fundamental "
         "em escolas que informaram ter suspendido ou interrompido as aulas por motivo de segurança, em termos de "
         "violência, pelo menos uma vez nos 12 meses anteriores à pesquisa, por dependência administrativa da escola, "
         "com indicação do intervalo de confiança de 95%, segundo as capitais das Unidades da Federação - 2015")
T_COND = ("Tabela 9.3.1 - Percentual de escolares de 13 a 17 anos que usaram capacete dentre aqueles que andaram de "
          "motocicleta ou moto nos 30 dias anteriores à pesquisa, por sexo, com indicação do intervalo de confiança "
          "de 95%, segundo os grupos de idade e as Grandes Regiões - 2024")


@pytest.fixture(scope="module")
def vocab(cfg):
    return Vocab.from_config(cfg.vocabulary)


@pytest.fixture(scope="module")
def symbols(cfg):
    return frozenset(cfg.symbols)


def test_parse_title_2024_school_informant():
    t = parse_title(T2024, ["Fonte: IBGE, ... Pesquisa Nacional de Saúde do Escolar 2024",
                            "Nota: (1) Não são apresentados resultados para escolas que deixaram a questão sem resposta."])
    assert (t.table_id, t.year, t.footer_year) == ("17.2.1", 2024, 2024)
    assert t.population_code == "ESCOLARES_13_17_ANOS" and (t.age_min, t.age_max) == (13, 17)
    assert t.reference_period_code == "12_MESES"
    assert t.informant.startswith("Escola")
    assert t.territorial_scope == "region" and t.confidence_level == 95
    assert t.nonresponse_excluded and not t.multi_response


def test_parse_title_2009_population_restriction_and_scope():
    t = parse_title(T2009)
    assert t.year == 2009 and t.population_code == "ESCOLARES_9ANO_EF_MENORES_18"
    assert t.reference_period == "últimos 30 dias" and t.reference_period_code == "30_DIAS"
    assert t.territorial_scope == "capital"
    assert "inferido" in t.informant_method  # informante não é inventado silenciosamente


def test_parse_title_2015_sample_from_footer():
    t = parse_title(T2015, ["Fonte: IBGE, ..., Pesquisa Nacional de Saúde do Escolar, Amostra 1, 2015."])
    assert t.sample == "Amostra 1" and t.estimate_type == "total_e_percentual" and t.territorial_scope == "capital"


def test_parse_title_conditional_denominator():
    t = parse_title(T_COND, ["Nota: Exclusive os escolares que não andaram de moto/motocicleta nos últimos 30 dias."])
    assert t.conditional_on.startswith("andaram de motocicleta")
    assert t.denominator_note.startswith("Exclusive")


def test_school_informant_without_director_mention_2019():
    t = parse_title("Tabela 17.2.1 - Percentual de escolares de 13 a 17 anos em escolas que tiveram de suspender ou "
                    "interromper suas aulas por motivo de segurança em termos de violência alguma vez nos 12 meses "
                    "anteriores à pesquisa, segundo as Grandes Regiões - 2019")
    assert t.informant.startswith("Escola") and t.year == 2019


def test_missing_patterns_stay_pending():
    t = parse_title("Tabela 99.1 - Algo sem padrões conhecidos")
    assert t.year is None and t.population == "PENDENTE" and t.territorial_scope == "PENDENTE"


def test_period_code_equivalences():
    assert period_code("últimos 30 dias") == period_code("30 dias anteriores à pesquisa") == "30_DIAS"
    assert period_code("alguma vez na vida") == "VIDA"


def test_header_tree_from_merges(src_factory, build_ibge_xlsx, vocab, symbols):
    g = read_xlsx(src_factory(), build_ibge_xlsx())[0]
    title, rows = extract_title(g)
    assert rows == [0]
    lay = parse_header(g, rows, vocab, symbols)
    assert [b.estimate_col for b in lay.blocks] == [1, 4, 7]
    tot, homem, mulher = lay.blocks
    assert tot.dims == {} and tot.ci_lower_col == 2 and tot.ci_upper_col == 3 and tot.ci_level == 95
    assert homem.dims == {"sex": "Homem"} and mulher.dims == {"sex": "Mulher"}
    assert tot.statement.startswith("Percentual de escolares") and tot.statement_kind == "proportion"
    assert not lay.problems


def test_header_leaf_total_is_estimate_2009_style(src_factory, vocab, symbols):
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Tabela 8.3 - Percentual de escolares frequentando o 9º ano que fizeram algo - 2009"
    ws["B2"] = "Percentual de escolares (%)"
    ws.merge_cells("B2:G2")
    ws["B3"], ws["E3"] = "Total", "Sexo"
    ws.merge_cells("B3:D3")
    ws.merge_cells("E3:G3")
    ws["E4"] = "Feminino"
    ws.merge_cells("E4:G4")
    ws["B5"], ws["C5"], ws["E5"], ws["F5"] = "Total", "Intervalo de confiança de 95%", "Total", "Intervalo de confiança de 95%"
    ws.merge_cells("C5:D5")
    ws.merge_cells("F5:G5")
    for c in "CF":
        ws[f"{c}6"] = "Limite infe- rior"
    for c in "DG":
        ws[f"{c}6"] = "Limite supe- rior"
    ws.append(["Total", 10, 9, 11, 8, 7, 9])
    bio = io.BytesIO()
    wb.save(bio)
    g = read_xlsx(src_factory(), bio.getvalue())[0]
    lay = parse_header(g, [0], vocab, symbols)
    assert len(lay.blocks) == 2
    assert lay.blocks[1].dims == {"sex": "Mulher"}
    assert lay.blocks[1].dim_labels == {"sex": "Feminino"}  # rótulo original preservado


def test_header_asymmetric_flagged(src_factory, vocab, symbols):
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Tabela 8.1 - Percentual de algo - 2009"
    ws["B2"] = "Percentual de escolares (%)"
    ws.merge_cells("B2:G2")
    ws["B3"], ws["E3"] = "Local", "Dependência administrativa da escola"
    ws.merge_cells("B3:D3")
    ws.merge_cells("E3:G3")
    ws["B4"], ws["E4"] = "Escola", "Pública"
    ws.merge_cells("B4:D4")
    ws.merge_cells("E4:G4")
    for c0 in "BE":
        ws[f"{c0}5"] = "Total"
    for c in "CF":
        ws[f"{c}5"] = "Limite inferior"
    for c in "DG":
        ws[f"{c}5"] = "Limite superior"
    ws.append(["Total", 5, 4, 6, 9, 8, 10])
    bio = io.BytesIO()
    wb.save(bio)
    lay = parse_header(read_xlsx(src_factory(), bio.getvalue())[0], [0], vocab, symbols)
    assert any(p[0] == "HDR006" for p in lay.problems)


def test_unknown_column_is_reported(src_factory, vocab, symbols):
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Tabela 1.1 - Percentual de algo - 2024"
    ws["B2"], ws["C2"] = "Percentual", "Coluna misteriosa"
    ws.append(["Brasil", 10, 99])
    bio = io.BytesIO()
    wb.save(bio)
    lay = parse_header(read_xlsx(src_factory(), bio.getvalue())[0], [0], vocab, symbols)
    assert any(p[0] == "HDR001" for p in lay.problems)
    assert len(lay.blocks) == 1


def test_header_without_merges_is_inferred_and_flagged(src_factory, vocab, symbols):
    csv_text = ("Tabela 1.1 - Percentual de algo - 2024;;;;;;\n"
                "Território;Sexo;;;;;\n"
                ";Homem;;;Mulher;;\n"
                ";Percentual;Limite inferior;Limite superior;Percentual;Limite inferior;Limite superior\n"
                "Brasil;10,0;9,0;11,0;12,0;11,0;13,0\n")
    g = read_csv(src_factory("t.csv", "csv"), csv_text.encode())[0]
    lay = parse_header(g, [0], vocab, symbols)
    assert [b.dims for b in lay.blocks] == [{"sex": "Homem"}, {"sex": "Mulher"}]
    assert any(p[0] == "HDR005" for p in lay.problems)


def test_stub_sections_and_age_total(src_factory, build_ibge_xlsx, vocab, symbols):
    g = read_xlsx(src_factory(), build_ibge_xlsx())[0]
    lay = parse_header(g, [0], vocab, symbols)
    stub = parse_stub(g, lay.body_start, lay.blocks, vocab, symbols, (13, 17))
    data = [r for r in stub.rows if r.kind == "data"]
    assert [r.label for r in data] == ["Brasil", "Norte"]
    assert all(r.context == {} for r in data), "'13 a 17 anos' = universo -> Total"
    assert stub.footnote_start is not None


@pytest.mark.parametrize("label,expected", [
    ("Total", ("total", "Total")), ("13 a 15 anos", ("age_group", "13 a 15 anos")),
    ("16 e 17 anos", ("age_group", "16 e 17 anos")), ("Homem", ("sex", "Homem")),
    ("Privada", ("admin_dependency", "Privada")), ("Indígenas", ("unknown", "Indígenas")),
])
def test_group_label_classification(vocab, label, expected):
    assert classify_group_label(label, vocab, (13, 17)) == expected


@pytest.mark.parametrize("label,scope,code", [
    ("Brasil", "uf", "BR"), ("Centro-Oeste", "region", "CO"), ("Distrito Federal", "uf", "53"),
    ("Distrito Federal", "capital", "5300108"), ("Brasília", "capital", "5300108"), ("São Luis", "capital", "2111300"),
    ("Rio de Janeiro", "uf", "33"), ("Rio de Janeiro", "capital", "3304557"), ("Total", "capital", "CAPITAIS"),
])
def test_territory_resolution_is_scope_aware(cfg, label, scope, code):
    assert TerritoryResolver(cfg).resolve(label, scope).territory.code == code


def test_total_outside_capital_scope_is_not_resolved(cfg):
    assert TerritoryResolver(cfg).resolve("Total", "uf").territory is None


def test_section_row_with_repeated_values_is_a_section(src_factory, vocab, symbols):
    """PeNSE 2019, tabela 9.10.1: a linha '13 a 15 anos' repete os números da linha 'Brasil' seguinte."""
    import io

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Tabela 9.10.1 - Percentual de escolares de 13 a 17 anos que fizeram algo - 2019"
    ws["B2"], ws["C2"], ws["D2"] = "Percentual", "Limite inferior", "Limite superior"
    ws.append(["13 a 17 anos"])
    ws.append(["Brasil", 2.9, 2.7, 3.1])
    ws.append(["13 a 15 anos", 2.8, 2.5, 3.0])
    ws.append(["Brasil", 2.8, 2.5, 3.0])
    ws.append(["16 e 17 anos", 9.9, 9.0, 10.0])
    ws.append(["Brasil", 3.0, 2.6, 3.4])
    bio = io.BytesIO()
    wb.save(bio)
    g = read_xlsx(src_factory(), bio.getvalue())[0]
    lay = parse_header(g, [0], vocab, symbols)
    stub = parse_stub(g, lay.body_start, lay.blocks, vocab, symbols, (13, 17))
    data = [r for r in stub.rows if r.kind == "data"]
    assert [r.context for r in data] == [{}, {"age_group": "13 a 15 anos"}, {"age_group": "16 e 17 anos"}]
    codes = [p[0] for p in stub.problems]
    assert codes == ["STB004", "STB005"], "repetida = artefato informativo; diferente = revisão"
