"""Tabelas com vários anos na mesma aba: cada valor recebe o ano da sua coluna/seção.

Nenhum arquivo enviado tem esse formato (só há 2009, 2015 e 2024, cada um em tabelas
próprias); estes testes garantem que, se uma tabela 2009-2015 aparecer, os anos não sejam
fundidos num só nem lidos como categoria de resposta.
"""
import io

import openpyxl

from pense.adapters.base import SheetContext
from pense.adapters.registry import select_adapter
from pense.ingestion.readers import read_xlsx
from pense.parsing.header import Vocab
from pense.parsing.title import extract_footnotes, extract_title, parse_title
from pense.standardization.territories import TerritoryResolver

TITLE = ("Tabela 9.9 - Percentual de escolares frequentando o 9º ano que se sentiram inseguros no trajeto "
         "casa-escola nos 30 dias anteriores à pesquisa, com indicação do intervalo de confiança de 95%, "
         "segundo as Grandes Regiões - 2009-2012")


def _ctx(cfg, grid):
    title, rows = extract_title(grid)
    notes, _ = extract_footnotes(grid, max(rows) + 1)
    info = parse_title(title, notes)
    info.title_rows = rows
    return SheetContext(grid, info, cfg, Vocab.from_config(cfg.vocabulary), frozenset(cfg.symbols),
                        TerritoryResolver(cfg))


def _year_columns_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "9.9"
    ws["A1"] = TITLE
    ws["B2"] = "Percentual de escolares (%)"
    ws.merge_cells("B2:G2")
    ws["B3"], ws["E3"] = 2009, 2012.0  # o Excel/xls guarda anos como números
    ws.merge_cells("B3:D3")
    ws.merge_cells("E3:G3")
    for c0, lo, hi in (("B", "C", "D"), ("E", "F", "G")):
        ws[f"{c0}4"], ws[f"{lo}4"], ws[f"{hi}4"] = "Percentual", "Limite inferior", "Limite superior"
    ws.append(["Brasil", 10.0, 9.0, 11.0, 14.0, 13.0, 15.0])
    ws.append(["Norte", 12.0, 10.0, 14.0, 16.0, 14.0, 18.0])
    ws.append([])
    ws.append(["Fonte: IBGE, Diretoria de Pesquisas, Pesquisa Nacional de Saúde do Escolar 2009-2012."])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_title_with_year_list_has_no_single_year():
    t = parse_title(TITLE)
    assert t.multi_year and t.years_in_title == [2009, 2012]
    assert t.year is None  # o título não decide o ano de nenhum valor
    assert "2009" not in t.statement


def test_year_columns_become_year_not_category(cfg, src_factory):
    ctx = _ctx(cfg, read_xlsx(src_factory(), _year_columns_workbook())[0])
    sel = select_adapter(ctx)
    assert sel.adapter is not None and sel.adapter.family_id == "PENSE_MULTIANO"
    out = sel.adapter.extract(ctx)
    by = {(r["territory_name"], r["edition_year"]): r for r in out.records}
    assert set(by) == {("Brasil", 2009), ("Brasil", 2012), ("Norte", 2009), ("Norte", 2012)}
    br12 = by[("Brasil", 2012)]
    assert (br12["estimate"], br12["ci_lower"], br12["ci_upper"], br12["source_cell"]) == (14.0, 13.0, 15.0, "E5")
    # série publicada como conjunto: id estável (não se mistura com a edição regular do ano) e ano por valor
    assert br12["edition_id"] == "PENSE_SERIE_9ANO" and br12["edition_year"] == 2012
    assert by[("Brasil", 2009)]["indicator_id"] == br12["indicator_id"] == "PENSE_SERIE_9ANO__9.9"
    assert all(r["response_category"] == "" for r in out.records), "ano não pode virar categoria"
    assert all("ANO_DA_CELULA" in r["extraction_flags"] for r in out.records)
    assert all(r["year_source"] == "cabeçalho/seção da tabela" for r in out.records)


def test_year_sections_in_stub(cfg, src_factory):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = TITLE
    ws["B2"] = "Percentual de escolares (%)"
    ws.merge_cells("B2:D2")
    ws["B3"], ws["C3"], ws["D3"] = "Percentual", "Limite inferior", "Limite superior"
    ws.append(["2009"])
    ws.append(["Brasil", 10.0, 9.0, 11.0])
    ws.append(["2012"])
    ws.append(["Brasil", 14.0, 13.0, 15.0])
    ws.append([])
    ws.append(["Fonte: IBGE, Diretoria de Pesquisas, Pesquisa Nacional de Saúde do Escolar 2009-2012."])
    bio = io.BytesIO()
    wb.save(bio)
    ctx = _ctx(cfg, read_xlsx(src_factory(), bio.getvalue())[0])
    out = select_adapter(ctx).adapter.extract(ctx)
    assert sorted((r["edition_year"], r["estimate"]) for r in out.records) == [(2009, 10.0), (2012, 14.0)]
    assert not [i for i in out.issues if i.rule.startswith("STB001")], "ano não é seção desconhecida"


def test_multi_year_value_without_year_is_flagged(cfg, src_factory):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = TITLE
    ws["B2"], ws["C2"], ws["D2"] = "Percentual", "Limite inferior", "Limite superior"
    ws.append(["Brasil", 10.0, 9.0, 11.0])
    ws.append([])
    ws.append(["Fonte: IBGE, Pesquisa Nacional de Saúde do Escolar 2009-2012."])
    bio = io.BytesIO()
    wb.save(bio)
    ctx = _ctx(cfg, read_xlsx(src_factory(), bio.getvalue())[0])
    out = select_adapter(ctx).adapter.extract(ctx)
    assert out.records and all("ANO_INDEFINIDO" in r["extraction_flags"] for r in out.records)
    assert all(r["edition_year"] is None for r in out.records)


def test_year_label_variants():
    from pense.parsing.title import year_label
    assert [year_label(x) for x in ("2009", "2009.0", "2015 (1)", "13 a 15 anos", "20090")] == \
        ["2009", "2009", "2015", "", ""]
