"""Leitores por formato que produzem ``Grid`` (valores + células mescladas).

* XLSX/XLSM: openpyxl (modo completo, necessário para ler mesclagens).
* XLS: xlrd com ``formatting_info=True`` (mesclagens nativas do BIFF).
* ODS: odfpy, lendo ``number-columns/rows-spanned`` como mesclagens.
* CSV: módulo csv, com detecção de codificação e delimitador; sem mesclagens.
"""
from __future__ import annotations

import csv
import io
from typing import Any, Callable

from ..models import SourceFile
from ..parsing.grid import Grid, Merge

MAX_ODS_REPEAT = 200  # colunas/linhas repetidas vazias ao fim de abas ODS


class ReadError(Exception):
    """Falha de leitura de um arquivo inteiro (registrada, não interrompe a pipeline)."""


def read_xlsx(src: SourceFile, data: bytes) -> list[Grid]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    grids: list[Grid] = []
    for idx, ws in enumerate(wb.worksheets):
        values = [[cell.value for cell in row] for row in ws.iter_rows()]
        merges = [Merge(r.min_row - 1, r.max_row - 1, r.min_col - 1, r.max_col - 1)
                  for r in ws.merged_cells.ranges]
        grids.append(Grid(src, ws.title, idx, values, merges, "native"))
    return grids


def read_xls(src: SourceFile, data: bytes) -> list[Grid]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    grids: list[Grid] = []
    for idx in range(book.nsheets):
        sh = book.sheet_by_index(idx)
        values: list[list[Any]] = []
        for r in range(sh.nrows):
            row: list[Any] = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(xlrd.error_text_from_code.get(cell.value, "#ERR"))
                else:
                    row.append(cell.value)
            values.append(row)
        merges = [Merge(r0, r1 - 1, c0, c1 - 1) for (r0, r1, c0, c1) in sh.merged_cells]
        grids.append(Grid(src, sh.name, idx, values, merges, "native"))
    return grids


def read_ods(src: SourceFile, data: bytes) -> list[Grid]:
    from odf import opendocument, table
    from odf import text as odftext
    from odf.namespaces import OFFICENS, TABLENS

    doc = opendocument.load(io.BytesIO(data))
    grids: list[Grid] = []
    for idx, tbl in enumerate(doc.spreadsheet.getElementsByType(table.Table)):
        name = tbl.getAttrNS(TABLENS, "name") or f"Sheet{idx + 1}"
        values: list[list[Any]] = []
        merges: list[Merge] = []
        for row_el in tbl.getElementsByType(table.TableRow):
            rrep = int(row_el.getAttrNS(TABLENS, "number-rows-repeated") or 1)
            row: list[Any] = []
            for cell in row_el.childNodes:
                if cell.qname[1] not in ("table-cell", "covered-table-cell"):
                    continue
                crep = int(cell.getAttrNS(TABLENS, "number-columns-repeated") or 1)
                vtype = cell.getAttrNS(OFFICENS, "value-type")
                if vtype in ("float", "percentage", "currency"):
                    val: Any = float(cell.getAttrNS(OFFICENS, "value"))
                else:
                    paras = cell.getElementsByType(odftext.P)
                    val = "\n".join(str(p) for p in paras) or None
                cspan = int(cell.getAttrNS(TABLENS, "number-columns-spanned") or 1)
                rspan = int(cell.getAttrNS(TABLENS, "number-rows-spanned") or 1)
                if cspan > 1 or rspan > 1:
                    r0, c0 = len(values), len(row)
                    merges.append(Merge(r0, r0 + rspan - 1, c0, c0 + cspan - 1))
                row.extend([val] * min(crep, MAX_ODS_REPEAT))
            for _ in range(min(rrep, MAX_ODS_REPEAT)):
                values.append(list(row))
        width = max((len(r) for r in values), default=0)
        values = [r + [None] * (width - len(r)) for r in values]
        grids.append(Grid(src, name, idx, values, merges, "native"))
    return grids


def read_csv(src: SourceFile, data: bytes) -> list[Grid]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 sempre decodifica
        raise ReadError("Codificação não reconhecida")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if text.count(";") > text.count(",") else ","
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    width = max((len(r) for r in rows), default=0)
    values = [[(v if v != "" else None) for v in r] + [None] * (width - len(r)) for r in rows]
    return [Grid(src, "CSV", 0, values, [], "unavailable")]


READERS: dict[str, Callable[[SourceFile, bytes], list[Grid]]] = {
    "xlsx": read_xlsx,
    "xlsm": read_xlsx,
    "xls": read_xls,
    "ods": read_ods,
    "csv": read_csv,
}


def read_grids(src: SourceFile, data: bytes) -> list[Grid]:
    reader = READERS.get(src.fmt)
    if reader is None:
        raise ReadError(f"Formato sem leitor: {src.fmt}")
    try:
        return reader(src, data)
    except ReadError:
        raise
    except Exception as exc:  # qualquer falha da biblioteca vira erro registrado
        raise ReadError(f"{type(exc).__name__}: {exc}") from exc
